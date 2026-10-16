"""Editable academic DOCX report, created from saved outputs only.

Use python-docx; plots/equation displays are embedded raster figures. Text,
headings, numeric tables, captions and repository details remain editable.
"""
from pathlib import Path
import csv,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT,WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from src.experiments.runner import ROOT,RESULTS_ROOT,CSV_DIR,TABLES_DIR
from src.project_config import (
    ACTIVE_PROVENANCE, ASSIGNMENT_ISSUED, OPTIMIZER_SEEDS, REPORT_TITLE,
    ROLL_SUFFIX, STUDENT, TASK_C_ALGORITHMS, TASK_C_ALGORITHM_NUMBERS,
)
FIGURES_DIR=ROOT/'figures'/ACTIVE_PROVENANCE
REPORT_PATH=ROOT/'Lab_Assignment_7_Sofia_Abidi_241071072.docx'

def read_csv(name):
    with (ROOT/name).open(encoding='utf-8') as f: return list(csv.DictReader(f))
def number(value):
    if value is None or str(value) in ('None','','nan','NaN'): return 'N/A'
    try: return f'{float(value):.4g}'
    except ValueError: return str(value)
def count(value):
    if value is None or str(value) in ('None','','nan','NaN'): return 'N/A'
    value=float(value)
    return f'{int(value):,}' if value.is_integer() else f'{value:,.4g}'
def ms(r,k):
    mean=number(r.get(k+'_mean')); sd=number(r.get(k+'_std'))
    return 'N/A' if mean=='N/A' else (mean if sd=='N/A' else mean+' ± '+sd)
def parameter_label(value):
    labels={'population':'Population N','inertia':'Inertia w','epsilon':'Clearing ε'}
    try: values=json.loads(value) if isinstance(value,str) else value
    except (TypeError,json.JSONDecodeError): return str(value)
    if not isinstance(values,dict): return str(value)
    return ', '.join(f'{labels.get(key,key.replace("_"," ").title())} = {number(item)}' for key,item in values.items())

def generate():
    content=json.loads((RESULTS_ROOT/'report_content.json').read_text(encoding='utf-8')); doc=Document()
    section=doc.sections[0]; section.page_width=Inches(8.27); section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.72); section.left_margin=section.right_margin=Inches(.78)
    section.header_distance=section.footer_distance=Inches(.32)
    section.different_first_page_header_footer=False
    doc.settings.odd_and_even_pages_header_footer=True
    borders=OxmlElement('w:pgBorders'); borders.set(qn('w:offsetFrom'),'page')
    borders.set(qn('w:display'),'allPages')
    for edge in ['top','left','bottom','right']:
        border=OxmlElement('w:'+edge); border.set(qn('w:val'),'single')
        border.set(qn('w:sz'),'6'); border.set(qn('w:space'),'18'); border.set(qn('w:color'),'808080')
        borders.append(border)
    # Word requires pgBorders before cols/docGrid in sectPr schema order.
    columns=section._sectPr.find(qn('w:cols'))
    section._sectPr.insert(section._sectPr.index(columns),borders)
    update_fields=OxmlElement('w:updateFields'); update_fields.set(qn('w:val'),'true')
    doc.settings.element.append(update_fields)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name]; style.font.name='Cambria'; style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Cambria')
    doc.styles['Normal'].font.size=Pt(11); doc.styles['Normal'].paragraph_format.line_spacing=1.08; doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(24); doc.styles['Title'].font.bold=True
    doc.styles['Subtitle'].font.size=Pt(13); doc.styles['Subtitle'].font.bold=True
    for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
        style=doc.styles[name]; style.font.size=Pt(size); style.font.bold=True
        style.paragraph_format.keep_with_next=True; style.paragraph_format.space_before=Pt(10); style.paragraph_format.space_after=Pt(5)
    doc.styles['Caption'].font.size=Pt(9); doc.styles['Caption'].font.italic=True
    toc_subsection=doc.styles.add_style('TOC Subsection',WD_STYLE_TYPE.PARAGRAPH)
    toc_subsection.base_style=doc.styles['Heading 2']
    toc_subsection.font.name='Cambria'; toc_subsection.font.size=Pt(13); toc_subsection.font.bold=True
    toc_subsection.paragraph_format.keep_with_next=True
    toc_subsection.paragraph_format.space_before=Pt(10); toc_subsection.paragraph_format.space_after=Pt(5)
    for header_part in (section.header,section.even_page_header):
        header=header_part.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER
        header_run=header.add_run('VEERMATA JIJABAI TECHNOLOGICAL INSTITUTE  |  OPTIMIZATION TECHNIQUES LAB')
        header_run.bold=True; header_run.font.name='Cambria'; header_run.font.size=Pt(8.5)
    for footer_part in (section.footer,section.even_page_footer):
        footer=footer_part.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(f'{STUDENT.name}  |  {STUDENT.registration}  |  Page ')
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.core_properties.title=REPORT_TITLE; doc.core_properties.subject='37 algorithms and Tasks A to F'; doc.core_properties.author=STUDENT.name
    def clean(text):
        text=str(text)
        replacements={'\u00e2\u20ac\u201d':'—','\u00e2\u20ac\u2122':'’','\u00e2\u2020\u2019':'→','\u00c2±':'±',' +/- ':' ± '}
        for old,new in replacements.items(): text=text.replace(old,new)
        text=text.replace('`','')
        text=re.sub(r'(?<![\d,])10000(?!\d)','10,000',text)
        text=text.replace('N/A ± N/A','N/A').replace('nan ± nan','N/A')
        return text
    def para(text,bold_label=None,left=False,style=None):
        text=clean(text)
        if bold_label: bold_label=clean(bold_label)
        p=doc.add_paragraph()
        if style: p.style=style
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.widow_control=True
        if bold_label: p.add_run(bold_label+' ').bold=True
        p.add_run(text); return p
    def heading(text,level=1,newpage=False,toc_sub=False):
        p=doc.add_heading(clean(text),level)
        if toc_sub: p.style=doc.styles['TOC Subsection']
        if newpage: p.paragraph_format.page_break_before=True
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        widths=widths or [6.77/len(headers)]*len(headers)
        for i,w in enumerate(widths): t.columns[i].width=Inches(w)
        for i,text in enumerate(headers): t.rows[0].cells[i].text=clean(text)
        for row in rows:
            cells=t.add_row().cells
            for i,value in enumerate(row): cells[i].text=clean(value)
        for ri,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr(); keep=OxmlElement('w:cantSplit'); pr.append(keep)
            if ri==0: repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side); el.set(qn('w:w'),'70'); el.set(qn('w:type'),'dxa'); margins.append(el)
                tcpr.append(margins); borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
                tcpr.append(borders)
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6EF' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF')); tcpr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2); p.paragraph_format.line_spacing=1.
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: run.font.name='Cambria'; run.font.size=Pt(8.6); run.bold=ri==0
        return t
    def toc():
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
        begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); begin.set(qn('w:dirty'),'true')
        instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' TOC \\o "1-1" \\t "TOC Subsection,2" \\h \\z '
        separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
        placeholder_text=OxmlElement('w:t'); placeholder_text.text='Update this field to build the table of contents.'
        end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
        for element in (begin,instr,separate,placeholder_text,end):
            run=OxmlElement('w:r'); run.append(element); p._p.append(run)
    def landscape_section():
        landscape=doc.add_section(WD_SECTION.NEW_PAGE)
        landscape.orientation=WD_ORIENT.LANDSCAPE
        landscape.page_width=Inches(11.69); landscape.page_height=Inches(8.27)
        landscape.top_margin=landscape.bottom_margin=Inches(.62)
        landscape.left_margin=landscape.right_margin=Inches(.62)
        landscape.header_distance=landscape.footer_distance=Inches(.28)
        return landscape
    def figure(name,caption,width=6.3):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(FIGURES_DIR/name),width=Inches(width))
        cap=doc.add_paragraph(caption,style='Caption'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_together=True
    def box(title,text):
        t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
        cell=t.cell(0,0); cell.text=''
        p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
        p.add_run(title.upper()+'\n').bold=True; p.add_run(text)
        tcpr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'8'); el.set(qn('w:color'),'666666'); borders.append(el)
        tcpr.append(borders); doc.add_paragraph().paragraph_format.space_after=Pt(2)
    def markdown(path,exclude_title=False,technical=False):
        for block in (ROOT/path).read_text(encoding='utf-8').split('\n\n'):
            block=block.strip()
            if not block: continue
            if block.startswith('#'):
                line=block.splitlines()[0]; hashes=len(line)-len(line.lstrip('#'))
                title=clean(line.lstrip('# ').strip())
                if not (exclude_title and hashes==1):
                    heading(title,min(hashes,3),toc_sub=technical and hashes==2)
                rest=' '.join(part.strip() for part in block.splitlines()[1:])
                if rest: para(rest,left=technical)
                continue
            lines=block.splitlines()
            if technical and any(re.match(r'^\d+\.\s',line.strip()) for line in lines):
                items=[]; current=[]
                for line in lines:
                    line=line.strip()
                    if re.match(r'^\d+\.\s',line):
                        if current: items.append(' '.join(current))
                        current=[line]
                    elif line: current.append(line)
                if current: items.append(' '.join(current))
                for item in items:
                    p=para(item,left=True); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18)
            else:
                para(' '.join(line.strip() for line in lines),left=technical)
    # Title page
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(STUDENT.institution.upper()); r.bold=True; r.font.name='Cambria'; r.font.size=Pt(16)
    p=doc.add_paragraph('B.TECH COMPUTER ENGINEERING',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('OPTIMIZATION TECHNIQUES LABORATORY',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('LAB ASSIGNMENT 7\nMULTIOBJECTIVE OPTIMIZATION',style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('37 Algorithms, Trade-offs and Applications',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    table(['Student detail','Value'],[['Name',STUDENT.name],['Registration number',STUDENT.registration],['Roll suffix',ROLL_SUFFIX],['Course',STUDENT.course],['Assignment issued',ASSIGNMENT_ISSUED],['Assigned deep dives',', '.join(f'A{n} {name}' for n,name in TASK_C_ALGORITHMS)]],widths=[2.0,4.2])
    box('Aim','To implement, verify and compare classical, evolutionary, swarm, indicator, preference and gradient-based multiobjective optimization methods on the specified P1-P14 problem bank.')
    box('Objectives','Implement all 37 assigned mechanisms; verify shared numerical utilities and constraints; study A35 R-NSGA-II, A10 VEGA and A22 MOPSO in depth; compare methods with reproducible metrics; and document suitable applications and limitations.')
    p=para('GitHub repository: https://github.com/sofiaabidi/OT_LAB07'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para('This report presents tested Python implementations of all 37 catalogue mechanisms, their measured tradeoffs on the assignment problems, and the three assigned study cards. It includes Tasks A to F, raw-result references and reproducible commands. The main finding is that selection geometry, constraints and preferences determine useful choices: no algorithm leads across every problem and metric.')
    para('Ten-trial stochastic studies use a 10000-evaluation cap, with P9 feature-selection experiments explicitly reduced to 200. Classical scalar solvers often consume fewer calls. Measured quality is separated from theoretical mechanism claims; some advantages and disadvantages needing direct experimental support still need targeted ablations.')
    contents_heading=doc.add_paragraph(); contents_heading.paragraph_format.page_break_before=True
    contents_heading.paragraph_format.keep_with_next=True
    contents_run=contents_heading.add_run('Contents'); contents_run.bold=True; contents_run.font.name='Cambria'; contents_run.font.size=Pt(16)
    titles=['1 Overview and completion status','2 Methodology and reproducibility','3 Problem formulations','4 Shared tools and verification','5 Tasks A and B algorithm implementations and comparison','6 Task C assigned study cards and hand traces','7 Task D head to head comparisons','8 Task E algorithm selection cases','9 Task F viva preparation','10 Findings and limitations','11 Reproducibility and repository instructions','12 References','13 Appendix evidence files and lecture alignment']
    toc()
    para('The paginated table above is a native Word field. It remains editable and can be refreshed after later edits with Update Field.')
    heading(titles[0])
    para(f"The project has 37 separately named algorithm modules, 37 defining-mechanism tests and 62 passing automated tests in total. The main manifest contains {content['rows_count']} optimizer experiment rows and 60 additional network learning curves. Every front figure is drawn from persisted arrays of decisions, objectives and violations, not mock data.")
    para(f'Roll assignment calculation for R={ROLL_SUFFIX}: '+f'(({ROLL_SUFFIX}-1) mod 37)+1={TASK_C_ALGORITHM_NUMBERS[0]}; '+f'(({ROLL_SUFFIX}+11) mod 37)+1={TASK_C_ALGORITHM_NUMBERS[1]}; '+f'(({ROLL_SUFFIX}+23) mod 37)+1={TASK_C_ALGORITHM_NUMBERS[2]}.')
    table(['Area','Verified outcome'],[['Implementations','All 37 modules have core-mechanism tests'],['Problem bank','P1 to P14 definitions plus NumPy network'],['Task A','Designated problem runs, figures and consolidated algorithm/problem implementation log'],['Task B','Four reasoned strengths and four risks per method, with evidence-linked analysis and explicit causal boundaries'],['Task C','Three study cards, deterministic traces, three settings x ten trials'],['Task D','P3 P5 P10 M5 P8 comparisons plus ranking and gradient studies'],['Task E','Six cases including constraints and five objectives'],['Task F','Eight official questions plus focused A10/A22/A35 code walkthrough preparation'],['Remaining','P9 at full suggested budget and further mechanism-specific causal ablations']],widths=[1.05,5.72])
    heading(titles[1])
    para('Implementations use NumPy and the standard library. Matplotlib creates plots. SciPy SLSQP is used only in algorithms 1 to 7 for permitted scalar subproblems; scalarization, reference/epsilon sweeps, deviation elimination and priorities are handwritten. Scikit-learn loads only the Wisconsin dataset; five-neighbour classification and five-fold evaluation are handwritten. Neural-network forward and backward passes use NumPy without autograd.')
    para(f'Stochastic optimizer seeds are {OPTIMIZER_SEEDS[0]} to {OPTIMIZER_SEEDS[-1]}, ten independent streams. Benchmark construction and Monte Carlo metric seeds are separate fixed configuration values. Population and output archive cap are 40 except methods whose archive naturally shrinks. Candidate evaluation budgets include local-search proposals. Objective-vector evaluation counts one candidate irrespective of objective count. Finite-difference solver probes count in classical methods, including objective evaluations invoked by constraints. Reported runtime is wall-clock and should not be treated as a machine-independent algorithm complexity.')
    para('All natural maximization objectives are negated before dominance. Feasible candidates beat infeasible ones; smaller normalized positive constraint residual wins among infeasible candidates. Feasibility threshold is 1e-10, with a 1e-8 numerical verification tolerance on scalar-solver outputs. This prevents roundoff from being confused with a materially infeasible design.')
    para('Quality normalization uses fixed problem constants listed in Section 3. Two-objective HV is an exact staircase area. Higher-dimensional HV uses 8192 identical uniform samples per comparison, seed 991, in [0,1.1]^M. Its absolute Monte Carlo uncertainty is bounded by box_volume/(2*sqrt(8192)), about 0.00735 for M=3 and 0.00890 for M=5; differences smaller than that require cautious interpretation. SD across ten optimizer seeds is distinct from HV sampling error. IGD uses 301 exact-front samples where an analytic front is available; no true-front IGD is invented for P4 P8 or P9. Spacing is the sample SD of nearest L1 distances; a singleton scores zero without offering coverage.')
    para('Early nonelitist algorithms use a passive recording archive for comparable outputs. The final breeding population is saved separately in NPZ files. Thus a good recorded front does not prove that VEGA MOGA NSGA or NPGA preserved its parents. Classical algorithms use a common cap but naturally different consumed counts, so these are equal-cap rather than equal-consumed-budget studies.')
    para('The five supplied lecture decks were checked against the code. The teaching NPGA, MOSA, IBEA, RVEA and HypE conventions are stated in the appendix. Algorithms changed during that check were rerun; current figures and tables use refreshed rows.')
    heading(titles[2])
    for spec in content['formulations']:
        heading(spec['key']+' '+spec['title'],2)
        para(spec['variables'],bold_label='Decision variables and encoding')
        figure('equation_'+spec['key']+'.png','Standard form objectives in minimization convention',6.25)
        para(spec['formula'],bold_label='Formula in code notation')
        para(spec['constraints'],bold_label='Constraints and handling')
        para(spec['interpretation'])
        para(spec['sanity_check'],bold_label='Numerical sanity check')
        para('Normalized objectives are (F-ideal)/(nadir-ideal). Fixed ideal='+str(spec['ideal'])+'; fixed nadir='+str(spec['nadir'])+'. Quality HV reference='+str(spec['reference'])+' in normalized coordinates. These are declared scaling constants; conservative lower bounds are used where the attainable ideal is unavailable.')
        para(spec['termination'],bold_label='Termination')
        relevant=[r for r in content['summary'] if r['problem']==spec['key'] and r['group']=='base']
        para('Evaluated methods: '+(', '.join(sorted(set(r['algorithm'] for r in relevant))) if relevant else ('TOPSIS and VIKOR on supplier data' if spec['key']=='P12' else 'See supplementary investigations'))+'. See the corresponding Task A pages and raw files.')
        if spec['key']=='P6':
            table(['Point','Decision','f1','f2','Circle lhs','CV'],[['A','(0,0)',0,50,25,0],['B','(1,1)',8,32,17,0],['C','(0,3)',36,29,34,.36],['D','(.2,3)',36.16,27.04,32.04,.2816]])
            para('Manual Deb comparison: A and B are feasible and trade off objectives; both outrank C and D. D outranks C because .2816 < .36. The layers are [A,B], [D], [C]. The second constraint is inactive for all four points; infeasible objective values do not override violation order.')
        if spec['key']=='P9':
            freq=sorted(read_csv(CSV_DIR/'feature_frequencies.csv'),key=lambda r:-float(r['front_frequency']))[:10]
            para('Most frequent features in the pooled retained feature-selection fronts are listed below. Frequencies pool the two methods and ten seeds; the reduced budget does not establish globally optimal masks.')
            table(['Feature','Front frequency'],[[r['feature'],number(r['front_frequency'])] for r in freq],widths=[4.8,1.97])
        if spec['key']=='P12': para('All five supplied suppliers are nondominated under cost minimization and quality delivery reliability directions. Ranking imposes an additional preference order on these tradeoffs.')
        if spec['key']=='P8': table(['City','x','y'],[[i,number(x),number(y)] for i,(x,y) in enumerate(json.loads((ROOT/'docs/p8_coords.json').read_text(encoding='utf-8')))],widths=[.7,3.035,3.035])
    heading(titles[3])
    para('moo_utils.py centralizes dominance, pairwise front sorting, crowding distance, archive filtering, normalization, positive-residual violation, exact 2D HV, Monte Carlo HV, IGD, Spacing and reproducible reference directions. Pairwise arrays expose the calculations directly and cost O(M*N^2); vectorization improves execution time without changing asymptotic work.')
    table(['Check','Input','Expected'],[['Dominance','(1,2) versus (2,3)','First dominates'],['Tradeoff','(1,3) versus (2,2)','Neither dominates'],['HV','(1,3) (2,2) (3,1); reference (4,4)','6'],['IGD identical fronts','same point sets','0'],['Crowding','(0,2) (1,1) (2,0)','infinity 2 infinity'],['Lecture car front','A B C D E F table','A B C D'],['Lecture Deb layers','A D feasible; B CV2; C CV5','A,D then B then C'],['DTLZ2 g0','tail variables all .5','squared objective sum 1'],['Network derivatives','central finite difference +/-1e-6','analytical match to 5 decimal places']],widths=[1.3,3.5,1.97])
    para('Final automated result: 62 tests passed. Tests include seed reproducibility for all 26 stochastic modules and bounds for recorded solutions. For P8 the crossover and both mutation variants are checked across 20 seeds; cities occur exactly once. See results/roll72/test_output.txt for the named checks. Tests verify defining mechanisms, not a global-optimality guarantee.')
    heading(titles[4])
    para('The consolidated Task A log below is generated from verified base summaries. Each algorithm/problem combination has its own row so budgets and metrics cannot be misaligned. Ranking methods report N/A for optimizer metrics and untimed runtime because those quantities were not recorded; no values are invented.')
    implementation_log=read_csv(RESULTS_ROOT/'implementation_log.csv')
    table(['No.','Algorithm','Problem','Evaluations / cap','HV','IGD','Runtime s','Status'],
          [[r['algorithm_number'],r['algorithm'],r['problems_used'],r['evaluation_budget'],r['HV'],r['IGD'],r['runtime_seconds'],r['status']] for r in implementation_log],
          widths=[.32,1.05,.65,1.05,1.0,1.0,.85,.85])
    para('Each following algorithm section uses the same sequence: core mechanism, implementation and parameters, verified result table, recorded plot, concise interpretation, then the four mechanism-based advantages and disadvantages required by Task B. Evidence-linked statements describe what these runs show without converting metric correlation into a causal proof.')
    para(content['family_discussion'])
    for alg in content['algorithms']:
        a=alg['algorithm_number']; heading(f'{a} {alg["algorithm"]}',2)
        para(alg['what_it_does'],bold_label='Core mechanism')
        para('Module '+alg['module']+'; implementation status '+alg['status']+'; defining mechanism test test_'+f'{a:02d}'+'_mechanism passes. '+alg['docstring'].replace('Lecture decks supplied on 12 October 2026; teaching variants documented in docs/lecture_alignment.md.','').replace('\n',' '),bold_label='Implementation and parameters')
        if alg['results']:
            table(['Problem','Trials','HV','IGD','Spacing','Runtime s','Evaluations / cap'],
                  [[r['problem']+(' M'+str(r['objectives']) if r['problem']=='P10' else ''),r['trials'],ms(r,'HV'),ms(r,'IGD'),ms(r,'Spacing'),ms(r,'runtime'),count(r['evaluations_mean'])+' / '+count(r['budget_cap'])] for r in alg['results']],
                  widths=[.7,.4,1.1,1.1,1.1,1.15,1.22])
        figure(alg['figure'],f'Figure A{a:02d} Recorded output on designated problems. Plotted seed and actual calls appear in the panel. The full ten-trial metrics are in the table.',6.25)
        para(alg['evidence_supported_strengths'],bold_label='Interpretation—empirical performance')
        para(alg['evidence_supported_limitations'],bold_label='Interpretation—measured limitations')
        para('; '.join(f'{i+1}. {x}' for i,x in enumerate(alg['advantages'].split(';'))),bold_label='Four advantages')
        para('; '.join(f'{i+1}. {x}' for i,x in enumerate(alg['disadvantages'].split(';'))),bold_label='Four disadvantages')
        para(alg['evidence_scope'])
        para(alg['suitable']+'. Avoid: '+alg['avoid']+'.',bold_label='Best use and avoid')
    heading(titles[5])
    markdown(Path('docs/study_cards.md'),exclude_title=True,technical=True)
    task_para=lambda text,bold_label=None: para(text,bold_label=bold_label,left=True)
    for a in TASK_C_ALGORITHM_NUMBERS:
        heading(f'Flow diagram A{a} {content["algorithms"][a-1]["algorithm"]}',3)
        figure(f'flow_A{a:02d}.png',f'Figure C{a} Current implementation control flow',5.6)
    traces=json.loads((CSV_DIR/'hand_traces.json').read_text(encoding='utf-8'))
    heading('R-NSGA-II two selection iterations',2)
    task_para('Tiny instance: P1 with f1=(x-3)^2 and f2=(x-7)^2, population N=4, natural-unit reference r=(4,4), fixed ideal=(0,0), nadir=(16,16), and therefore normalized reference (0.25,0.25). Epsilon=0.20. Parents and four fixed offspring form eight candidates. Deb nondominated sorting is applied first; within the splitting front, candidates are ordered by Euclidean reference distance and epsilon clearing controls local spacing.')
    for r in traces['r_nsga2']:
        heading(f'R-NSGA-II iteration {r["iteration"]}',3)
        task_para(f"Parents={r['parents_x']}; offspring={r['offspring_x']}; combined indices 0-3 are parents and 4-7 are offspring.")
        table(['i','Origin','x','f1','f2','z1','z2','d(i,r)'],[[i,'P' if i<4 else 'O',number(x),number(r['F'][i][0]),number(r['F'][i][1]),number(r['normalized_F'][i][0]),number(r['normalized_F'][i][1]),number(r['distance_to_reference'][i])] for i,x in enumerate(r['candidates_x'])],widths=[.3,.48,.55,.75,.75,.8,.8,1.25])
        task_para('Nondominated fronts (zero-based candidate indices): '+str(r['fronts'])+'. All eight candidates trade off f1 against f2, so the first front is also the splitting front.')
        table(['Order','i','x','Reference d','Min separation','Decision'],[[order+1,d['index'],number(r['candidates_x'][d['index']]),number(d['reference_distance']),number(d['minimum_selected_separation']),d['decision']] for order,d in enumerate(r['clearing_decisions'])],widths=[.5,.35,.7,1.1,1.2,2.4])
        task_para(f"Selected survivor indices={r['selected_indices']}, hence x={r['selected_x']}. These four survivors become the parents shown in the next iteration.")
    task_para('Code verification: tests/test_traces.py recomputes normalized objectives and Euclidean reference distances, then asserts the recorded clearing decisions and survivor indices for both iterations.')
    heading('VEGA two objective-wise selections',2)
    task_para('Tiny instance: the same P1 minimization problem and population N=4. Mating slots 1 and 3 use f1; slots 2 and 4 use f2. Each slot has an explicit two-candidate tournament. Variation uses the implementation\'s SBX equation c1=0.5[(1+beta)p1+(1-beta)p2], c2=0.5[(1-beta)p1+(1+beta)p2], with u=0.25, beta=(2u)^(1/16), positive sign, and zero mutation increment.')
    for r in traces['vega']:
        heading(f'VEGA iteration {r["iteration"]}',3)
        table(['i','x','f1 min','f2 min'],[[i,number(x),number(r['F'][i][0]),number(r['F'][i][1])] for i,x in enumerate(r['population'])],widths=[.5,1.1,2.2,2.2])
        table(['Slot','Objective','Tournament indices','Candidate x values','Winner','Winner x'],[[slot+1,'f'+str(obj+1),str(pair),str([number(r['population'][i]) for i in pair]),r['winner_indices'][slot],number(r['mating_pool'][slot])] for slot,(obj,pair) in enumerate(zip(r['slot_objectives'],r['tournament_pairs']))],widths=[.45,.65,1.2,1.75,.65,1.0])
        pool='['+', '.join(number(x) for x in r['mating_pool'])+']'
        pairs='['+', '.join('['+', '.join(number(x) for x in pair)+']' for pair in r['parent_pairs'])+']'
        increments='['+', '.join(number(x) for x in r['mutation_increments'])+']'
        task_para(f"Objective-wise mating pool={pool}; parent pairs={pairs}. With u={number(r['sbx_u'])}, beta={number(r['sbx_beta'])}; mutation increments={increments}.")
        table(['Child','x after SBX','f1','f2','Next population'],[[i+1,number(x),number(r['offspring_F'][i][0]),number(r['offspring_F'][i][1]),number(r['next_population'][i])] for i,x in enumerate(r['offspring'])],widths=[.65,1.35,1.45,1.45,1.45])
        task_para('The offspring replace the parents without elitist parent-plus-child survival. This explicit update produces the population used at the start of the next iteration.')
    task_para('Code verification: tests/test_traces.py checks the objective-wise winner indices and independently recomputes every parent and offspring objective used in both population updates.')
    heading('MOPSO two velocity iterations',2)
    task_para('Tiny instance: this is an independent scalar MOPSO teaching problem, not assignment P2. It uses x in [1,20], teaching velocity limits [-10,10], and objectives f1=100/x and f2=x^2, both minimized. These fixed teaching limits deliberately differ from the main implementation default, which caps each velocity component at 20% of its decision-variable range. The wider toy domain makes x=12 and x=15 valid. The update is v-new=w v+c1 r1(pbest-x)+c2 r2(leader-x), followed by clipping and x-new=x+v-new. Here w=0.7, c1=c2=1.5, and every r1=r2=0.5. Leaders are explicit archive members. A personal best changes only if the new position dominates it.')
    for r in traces['mopso']:
        heading(f'MOPSO iteration {r["iteration"]}',3)
        task_para(f"Archive before update={r['archive_before']}; selected leaders={r['leader']}; coefficients w={r['inertia']}, c1={r['c1']}, c2={r['c2']}, r1=r2=0.5.")
        table(['i','x','v','pbest','leader','wv','cognitive','social','v-new','x-new'],[[i,number(r['old_x'][i]),number(r['old_velocity'][i]),number(r['pbest'][i]),number(r['leader'][i]),number(r['inertia_term'][i]),number(r['cognitive_term'][i]),number(r['social_term'][i]),number(r['new_velocity'][i]),number(r['new_x'][i])] for i in range(len(r['new_x']))],widths=[.25,.55,.55,.62,.62,.55,.75,.65,.62,.62])
        table(['i','Old objectives','New objectives','pbest decision','pbest after'],[[i,str([number(v) for v in r['old_F'][i]]),str([number(v) for v in r['F'][i]]),r['pbest_decisions'][i],number(r['pbest_after'][i])] for i in range(len(r['new_x']))],widths=[.35,1.4,1.4,2.35,1.0])
        task_para(f"Every move improves one objective and worsens the other, so each trial is incomparable with its pbest and the pbest values remain {r['pbest_after']}. The nondominated archive becomes {r['archive_after']} (size {r['archive_size_after']}).")
    task_para('Code verification: tests/test_traces.py checks each inertia, cognitive and social term; verifies x-new=x+v-new; recomputes objectives; validates pbest decisions; and confirms that each updated archive is pairwise nondominated.')
    heading('Ten trial sensitivity studies',2)
    for a in TASK_C_ALGORITHM_NUMBERS:
        selected=[r for r in content['summary'] if r['group']=='sensitivity' and r['algorithm_number']==a]
        heading(content['algorithms'][a-1]['algorithm'],3)
        table(['Setting','Trials','HV','IGD','Spacing','Evaluations'],[[parameter_label(r['parameters']),r['trials'],ms(r,'HV'),ms(r,'IGD'),ms(r,'Spacing'),count(r['evaluations_mean'])] for r in selected],widths=[1.6,.45,1.2,1.2,1.2,1.12])
        best=max(selected,key=lambda r:r['HV_mean']); worst=min(selected,key=lambda r:r['HV_mean'])
        mechanism={35:'Epsilon controls the minimum normalized separation near the reference: increasing it strengthens local clearing but can force distance-based backfill.',10:'Population size trades simultaneous coverage against the number of generations available under the same 10000-evaluation budget.',22:'Inertia controls retained momentum: low values emphasize current leaders, while high values explore more aggressively and may overshoot.'}[a]
        caveat=(' P4 has no exact reference front, so IGD is intentionally unavailable and HV/Spacing carry the interpretation.' if a==22 else '')
        task_para(mechanism+' Largest mean HV occurred at '+parameter_label(best['parameters'])+' ('+ms(best,'HV')+'); smallest at '+parameter_label(worst['parameters'])+' ('+ms(worst,'HV')+').'+caveat+' These three settings and ten trials are a finite sensitivity comparison, not a claim of a globally optimal parameter. Overlapping sample SD requires caution.')
        figure(f'C_A{a:02d}.png',f'Figure S{a} Sensitivity metrics with sample SD across ten seeds',6.25)
    heading(titles[6])
    para('The continuous comparison includes 25 applicable two-objective optimizer mechanisms. Ranking methods are evaluated separately; ant construction requires tours, gradient methods require differentiability, and many-objective-specific methods are compared on P10. Algorithms 5 to 7 retain their designated utility goals and priorities rather than receiving invented P3/P5 goals. Algorithms 1 to 4 use at least 21 sweep settings, with metric p variants combined.')
    for p,m in [('P3',2),('P5',2),('P10',5),('P8',2)]:
        heading(p+' comparison',2)
        selected=sorted([r for r in content['summary'] if r['group']=='D' and r['problem']==p and r['objectives']==m],key=lambda r:-r['HV_mean'])
        table(['Rank','Algorithm','HV','IGD','Spacing','Runtime s'],[[i+1,r['algorithm'],ms(r,'HV'),ms(r,'IGD'),ms(r,'Spacing'),ms(r,'runtime')] for i,r in enumerate(selected)],widths=[.4,1.7,1.15,1.15,1.15,1.22])
        best=selected[0]; low=selected[-1]
        para('Ranked by mean HV: '+best['algorithm']+' leads at '+ms(best,'HV')+', compared with '+low['algorithm']+' at '+ms(low,'HV')+'. HV and IGD judge different aspects; this ranking is specific to the fixed scales and reference. All rows have ten trials. '+('Classical actual calls differ despite the common 10000 cap; see raw_experiments.csv.' if p in ['P3','P5'] else 'All candidate budgets are 10000.'))
        projection_note=(' The overlaid parallel-coordinate lines show every recorded five-objective front vector; this projection is descriptive rather than a dominance ordering.' if p=='P10' else ' The plot shows the recorded nondominated objective vectors for the stated problem and seed.')
        figure('D_'+p+'.png',f'Figure D{p} Combined seed {OPTIMIZER_SEEDS[0]} results.'+projection_note,6.25)
        if p=='P10':
            pressure=read_csv(CSV_DIR/'dominance_pressure.csv')
            for M in [3,5]:
                z=[float(r['nondominated_fraction']) for r in pressure if int(r['M'])==M]; para('M='+str(M)+' random population nondominated fraction mean='+number(sum(z)/len(z))+'. Increasing nondomination reduces rank discrimination, while vector association adds a directional criterion.')
        if p=='P8':
            bases=read_csv(CSV_DIR/'routing_baselines.csv'); table(['Nearest neighbour','Distance','Time','Tour'],[[r['baseline'],number(r['distance']),number(r['time']),r['tour']] for r in bases],widths=[1.1,.9,.9,3.87])
            para('Distance/time baselines use different edge metrics and retain valid 15-city permutations. A good optimized archive can improve one objective while worsening the other, so neither baseline should be assumed to dominate the entire tradeoff set.')
    heading('P3 population and generation budget investigation',2)
    population=read_csv(TABLES_DIR/'population_study_summary.csv')
    table(['Population','Budget','Full generations','HV','IGD','Spacing'],[[r['population'],r['budget'],r['full_generations'],ms(r,'HV'),ms(r,'IGD'),ms(r,'Spacing')] for r in population],widths=[.75,.7,1.1,1.4,1.4,1.42])
    para('This additional NSGA-II study uses ten seeds for each of nine population and budget settings. At fixed budget, larger populations receive fewer generations; improving population coverage can compete with convergence. More evaluations generally allow more refinement, but the measured IGD values determine the conclusion for each setting rather than an assumed universal trend. Ninety raw fronts and objective metrics are stored under results/roll72/csv/population/.')
    heading('Ranking existing alternatives and preferred scalar solutions',2)
    rankings=read_csv(CSV_DIR/'rankings.csv')
    table(['Input','Method','Weights','v','Top three'],[[r['problem'],r['method'],r['weights'],r['v'],r['top_three']] for r in rankings],widths=[1.2,.7,1.5,.4,2.97])
    para('P12 supplier identifiers are 1-based. P4 front identifiers are 0-based rows in P4_nsga2_ranking_input.csv. Moving .1 weight to the first criterion and subtracting .1 from the second preserves the simplex. VIKOR v=.25,.5,.75 and its compromise-set flags remain in rankings.csv.')
    for method in ['topsis','vikor']:
        subset=[r for r in rankings if r['problem']=='P12' and r['method']==method and float(r['v'])==.5]
        para(method.upper()+' top-three order '+('changes' if subset[0]['top_three']!=subset[1]['top_three'] else 'is unchanged')+' under the .1 weight shift: '+subset[0]['top_three']+' -> '+subset[1]['top_three']+'. This measured local response is not universal preference stability.')
    para('Value Function returns declared portfolio profiles or the laptop aspiration. Goal Programming uses production targets 2400 profit and 600 units, and normalized versus unit deviations. Lexicographic profit-first then units selects the profit tie with maximum units, approximately (300,200). Exact decisions and solver success messages are retained in the classical job JSON records; a solver failure is not silently treated as proof of optimality.')
    heading('P13 gradient methods and equal weights',2)
    table(['Method','Scale','BCE initial','BCE final','MSE initial','MSE final'],[[r['method'],r['regression_scale'],number(r['classification_initial']),number(r['classification_final_mean'])+' ± '+number(r['classification_final_std']),number(r['regression_initial']),number(r['regression_final_mean'])+' ± '+number(r['regression_final_std'])] for r in content['network']],widths=[.85,.45,1.1,1.65,1.1,1.62])
    para('The network has a five-input shared tanh layer of width eight and separate classification and regression heads. Ten seeds vary initialization while synthetic observations remain fixed. Regression target scale 10 creates larger gradient norms. Equal weights and MGDA sum to one; GradNorm weights sum to two by definition. Their model learning rates are explicitly recorded, so weight normalization changes effective step size and must be considered when interpreting loss differences.')
    for scale in [1,10]: figure(f'network_scale{scale}.png',f'Figure N{scale} Network task losses with mean and SD from ten initializations',6.25)
    para('On the toy, fixed equal weights converge to (0,0), while MGDA generally stops at a different Pareto-stationary point and GradNorm changes the task weighting. Individual losses need not both approach zero: the two target points conflict. The saved 10000-step histories and toy_equal.csv expose this trajectory rather than claiming full-front coverage from one initial point.')
    heading('Toy task losses against equal weights',2)
    figure('toy_comparison.png','Figure T13 Hand-coded toy loss trajectories against fixed equal weights, with identical starting point and learning rate',6.25)
    para('The equal-weight toy baseline tends to x=(0,0), giving losses (18,18). MGDA tends to a different point on the Pareto segment, while GradNorm adapts its loss weights. A smaller task 1 loss can coexist with a larger task 2 loss at a Pareto-stationary tradeoff; neither method is expected to make both target distances zero.')
    heading(titles[7])
    for c in content['cases']:
        heading(c['problem'],2); para('Suitable: '+c['suitable_algorithms']+'. Poor choice: '+c['poor_algorithm']+'. '+c['reason']); para(c['evidence']); para(c['interpretation'])
    heading(titles[8])
    for i,(q,answer) in enumerate(content['questions']):
        heading(str(i+1)+' '+q,3); para(answer)
    heading('Code Walkthrough Preparation',2)
    for name,answer in content['code_walkthrough']:
        heading(name,3); para(answer)
    heading(titles[9])
    para('The numerical comparison shows meaningful differences in convergence, coverage and computation across methods, but finite outputs do not establish global optima. Weighted Sum on P3 illustrates a structural unsupported-region limit. Many-objective methods preserve directional niches when ranks tie. Routing methods require valid edge/tour operations. Preference and gradient methods should be judged by their declared decision or learning task rather than by full-front HV alone.')
    para('Remaining work is explicit: P9 uses 200 evaluations per run rather than the suggested 10000, and its smallest observed feature set is an exploratory result. The Task B table contains four argued advantages and disadvantages plus traceable measured metrics, but targeted ablations do not yet establish two experiment-supported benefits and drawbacks for every method. This report does not claim that stricter evidence requirement is fully satisfied. Student explanation and personal conclusions remain necessary for the viva.')
    investigations=read_csv(CSV_DIR/'problem_investigations.csv'); table(['Problem','Investigation','Measured value','Interpretation'],[[r['problem'],r['investigation'],number(r['value']),r['notes']] for r in investigations],widths=[.55,2.15,.9,3.17])
    para('Additional limitations: fixed population 40 limits front resolution; sparse grids and reference lattices can bias irregular fronts; Monte Carlo selection noise is larger than final HV evaluation noise; runtimes measured in a process pool are environment specific; IGD depends on reference sample density; exact fronts for P4 P8 P9 remain unknown here. The shared output archive can mask nonelitist final-population losses, so saved final_f is essential for diagnosing collapse.')
    heading(titles[10])
    para('Keep src tests results figures docs and data together. From the project root create a Python 3.10+ virtual environment, install requirements.txt, run python -m unittest discover -s tests -v, and execute python run_all.py experiments --workers 4. Jobs checkpoint per algorithm/problem/seed; an existing JSON row is reused unless --force is requested. Use python run_all.py report to rebuild figures and the Word document from saved arrays.')
    para('GitHub repository: https://github.com/sofiaabidi/OT_LAB07. The README gives individual algorithm commands and selective reruns. This editable DOCX is the requested report deliverable; no PDF or submission ZIP was generated.')
    heading(titles[11]); markdown(Path('docs/references.md'),exclude_title=True)
    heading(titles[12])
    table(['Evidence','Location'],[['Raw fronts and decisions','results/roll72/csv/job_name.npz'],['Job metrics and scalar solver messages','results/roll72/csv/job_name.json'],['All raw experiment rows','results/roll72/csv/raw_experiments.csv'],['Mean and sample SD tables','results/roll72/tables/summary.csv'],['37 method comparison','results/roll72/tables/comparison.csv'],['Implementation log','results/roll72/implementation_log.csv'],['Deterministic full traces','results/roll72/csv/hand_traces.json'],['Feature frequencies','results/roll72/csv/feature_frequencies.csv'],['Network loss histories','results/roll72/csv/network_method_scale_seed.csv'],['Figure provenance','results/roll72/tables/figure_manifest.csv'],['Tests','results/roll72/test_output.txt'],['Requirement map','docs/requirements_checklist.md']],widths=[2.6,4.17])
    markdown(Path('docs/lecture_alignment.md'),exclude_title=True)
    landscape_section()
    heading('Consolidated Task B comparison table',2)
    para('The required six fields are grouped by algorithm family. Advantages and disadvantages are mechanism-based arguments; the detailed algorithm pages separately connect two observed strengths and two observed limitations to stored numerical evidence wherever justified.')
    families=[]
    for algorithm in content['algorithms']:
        if algorithm['family'] not in families: families.append(algorithm['family'])
    for family_name in families:
        heading(family_name,3)
        members=[a for a in content['algorithms'] if a['family']==family_name]
        table(['Algorithm','Main mechanism','Four advantages','Four disadvantages','Suitable','Avoid'],[[str(a['algorithm_number'])+' '+a['algorithm'],a['what_it_does'],'; '.join(str(i+1)+'. '+x for i,x in enumerate(a['advantages'].split(';'))),'; '.join(str(i+1)+'. '+x for i,x in enumerate(a['disadvantages'].split(';'))),a['suitable'],a['avoid']] for a in members],widths=[1.1,1.55,2.25,2.25,1.35,1.35])
    # python-docx's bundled template can carry paragraph borders; strip them.
    for tree in [doc._element,doc.styles.element]:
        for border in list(tree.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.save(REPORT_PATH)
    print(f'Saved editable report: {REPORT_PATH.name}')
if __name__=='__main__': generate()
